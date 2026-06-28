from __future__ import annotations

from pathlib import Path


def load_text(path: str | Path) -> list[dict]:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".txt", ".md", ".py", ".c", ".cpp", ".h", ".json", ".yaml", ".csv"}:
        return [{"page_number": None, "text": p.read_text(encoding="utf-8", errors="ignore")}]
    if suffix == ".pdf":
        try:
            import fitz
            doc = fitz.open(str(p))
            return [{"page_number": i + 1, "text": page.get_text("text")} for i, page in enumerate(doc)]
        except Exception:
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            return [{"page_number": i + 1, "text": page.extract_text() or ""} for i, page in enumerate(reader.pages)]
    if suffix == ".docx":
        import docx
        document = docx.Document(str(p))
        return [{"page_number": None, "text": "\n".join(par.text for par in document.paragraphs)}]
    if suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(p), data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None]
                if values:
                    lines.append(" | ".join(values))
        return [{"page_number": None, "text": "\n".join(lines)}]
    raise ValueError(f"暂不支持的文件类型：{suffix}")
